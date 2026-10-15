from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pypdf import PdfReader
from sqlalchemy import select

from app.api.deps import DbSession, StaffUser
from app.db.models import KnowledgeDocument
from app.schemas import KnowledgeCreate, KnowledgeRead
from app.services.rag import embed_text


router = APIRouter(prefix="/knowledge", tags=["knowledge"])


@router.get("", response_model=list[KnowledgeRead])
def list_documents(db: DbSession, _: StaffUser) -> list[KnowledgeDocument]:
    return list(db.scalars(select(KnowledgeDocument).order_by(KnowledgeDocument.created_at.desc())))


@router.post("", response_model=KnowledgeRead, status_code=201)
def create_document(payload: KnowledgeCreate, db: DbSession, _: StaffUser) -> KnowledgeDocument:
    document = KnowledgeDocument(
        **payload.model_dump(),
        embedding=embed_text(payload.content),
        chunk_count=_chunk_count(payload.content),
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


@router.post("/upload", response_model=KnowledgeRead, status_code=201)
async def upload_document(
    db: DbSession,
    _: StaffUser,
    file: UploadFile = File(...),
    title: str = Form(""),
) -> KnowledgeDocument:
    content = await file.read()
    text = _extract_text(file.filename or "document.txt", content)
    document = KnowledgeDocument(
        title=title or Path(file.filename or "知识文档").stem,
        source=file.filename or "upload",
        content=text,
        embedding=embed_text(text),
        chunk_count=_chunk_count(text),
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def _extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
    elif suffix == ".docx":
        text = "\n".join(paragraph.text for paragraph in Document(BytesIO(content)).paragraphs)
    elif suffix in {".txt", ".md"}:
        text = content.decode("utf-8-sig")
    else:
        raise HTTPException(status_code=400, detail="仅支持 PDF、DOCX、TXT 和 Markdown")
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 20:
        raise HTTPException(status_code=400, detail="文档没有可索引的文本")
    return text


def _chunk_count(text: str, chunk_size: int = 500) -> int:
    return max(1, (len(text) + chunk_size - 1) // chunk_size)
